import os
import datetime
import sqlite3
import json
from docx import Document
import PyPDF2
import openpyxl
import pytesseract
from PIL import Image

class FileExtractor:
    def __init__(self, file_full_path):
        self.full_path = file_full_path
        self.extract_metadata()

    def extract_text(self):
        file_full_path = self.full_path
        
        if os.path.basename(file_full_path).startswith("~$"):
            return ''

        if file_full_path.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp')):
            return self.extract_text_from_image(file_full_path)
        elif file_full_path.lower().endswith('.pdf'):
            return self.extract_text_from_pdf(file_full_path)
        elif file_full_path.lower().endswith('.docx'):
            return self.extract_text_from_docx(file_full_path)
        elif file_full_path.lower().endswith('.xlsx'):
            return self.extract_text_from_excel(file_full_path)
        else:
            return "Format de fichier non pris en charge"
    
    def extract_text_from_image(self, image_path):
        text = pytesseract.image_to_string(Image.open(image_path))
        return text

    def extract_text_from_pdf(self, pdf_path):
        with open(pdf_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page_num in range(len(pdf_reader.pages)):
                text += pdf_reader.pages[page_num].extract_text()
        return text

    def extract_text_from_docx(self, docx_path):
        doc = Document(docx_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text
        return text

    def extract_text_from_excel(self, xlsx_path):
        workbook = openpyxl.load_workbook(xlsx_path)
        sheet = workbook.active
        text = ""
        for row in sheet.iter_rows():
            for cell in row:
                text += str(cell.value) + " "
        return text

    def extract_metadata(self):
        self.file_name = os.path.basename(self.full_path)
        self.file_path = os.path.dirname(self.full_path)
        self.creation_date = datetime.datetime.fromtimestamp(os.path.getctime(self.full_path))
        
    def get_file_name(self):
        return self.file_name
    
    def get_full_path(self):
        return self.full_path
        
    def get_file_path(self):
        return self.file_path


class TextExtractor:
    def __init__(self, watched_folder, text_folder, excluded_folders=[], excluded_files=[]):
        self.watched_folder = watched_folder
        self.text_folder = text_folder
        self.excluded_folders = excluded_folders
        self.excluded_files = excluded_files
        self.keyword_extractor = ExtractKeyWord()

    def extract_and_save_text(self):
        for root_folder, _, files in os.walk(self.watched_folder):
            if self._should_exclude(root_folder):
                continue
            for file in files:
                if file not in self.excluded_files:
                    file_path = os.path.join(root_folder, file)
                    extracted_data = FileExtractor(file_path)
                    text = extracted_data.extract_text()
                    self._save_to_text_file(text, extracted_data)
                    if file.endswith(".txt"):
                        keywords = self.keyword_extractor.extraire_mot_cles(file_path)
                        self._update_keywords_in_database(file_path, keywords)

    def _should_exclude(self, folder):
        for excluded_folder in self.excluded_folders:
            if excluded_folder in folder:
                return True
        return False

    def _save_to_text_file(self, text, extracted_data):
        relative_path = os.path.relpath(extracted_data.get_file_path(), self.watched_folder)
        text_folder = os.path.join(self.text_folder, relative_path)
        os.makedirs(text_folder, exist_ok=True)
        text_file_name = os.path.splitext(extracted_data.get_file_name())[0] + ".txt"
        text_file_path = os.path.join(text_folder, text_file_name)
        with open(text_file_path, "w", encoding="utf-8") as text_file:
            text_file.write(text)

    def _update_keywords_in_database(self, file_path, keywords):
        conn = sqlite3.connect('metadata.db')
        c = conn.cursor()
        c.execute("UPDATE files SET keywords = ? WHERE file_path = ?", (json.dumps(keywords), file_path))
        conn.commit()
        conn.close()


class ExtractKeyWord:
    def __init__(self):
        self.exclusions = {
            "a", "z", "e", "r", "t", "y", "u", "i", "o", "p", "q", "s", "d", "f", "g", "h", "j", "k", "l", "m", "w", "x",
            "c", "v", "b", "n", "avec", "de", "du", "le", "la", "un", "une", "dans", "pour", "et", "après", "sur", "les",
            "ce", "en", "il", "je", "tu", "elle", "nous", "vous", "ils", "elles", "on", "Â«", "Â»", "Ã", "à"
        }
        # Ajouter même s'ils sont en majuscule
        self.exclusions = self.exclusions.union(set(map(str.upper, self.exclusions)))
        # Ajoute les caractères
        self.exclusions.update({',', ';', ':', '.', '!', '?', '(', ')', '[', ']', '{', '}', '<', '>', '/', '\\', '-',
                                '_', '=', '+', '*', '&', '%', '@', '#', '$'})

    def compter_mots_repetes(self, nom_fichier):
        mots_sans_exclusions = []
        with open(nom_fichier, 'r', encoding='utf-8') as file:
            data = file.read()
            words = data.split()
            for word in words:
                # Ignorer les mots vides et les mots dans les exclusions
                if word.strip() and word.lower() not in self.exclusions:
                    mots_sans_exclusions.append(word)

        # Créer un dictionnaire pour stocker le nombre de répétitions de chaque mot
        word_count = {}
        for word in mots_sans_exclusions:
            word_count[word] = word_count.get(word, 0) + 1

        # Convertir le dictionnaire en une liste de tuples (mot, nombre d'occurrences)
        mots_et_occurrences = list(word_count.items())
        # Trier la liste par nombre d'occurrences décroissant
        mots_et_occurrences.sort(key=lambda x: x[1], reverse=True)
        return mots_et_occurrences

    def extraire_mots_majuscules(self, nom_fichier):
        mots_majuscules = {}
        with open(nom_fichier, 'r', encoding='utf-8') as file:
            data = file.read()
            words = data.split()
            for word in words:
                # Ignorer les mots vides, les mots non entièrement en majuscules et les mots dans les exclusions
                if word.strip() and word.isupper() and word not in self.exclusions:
                    mots_majuscules[word] = mots_majuscules.get(word, 0) + 1

            # Convertir le dictionnaire en une liste de tuples (mot, nombre d'occurrences)
            mots_majuscules_tries = sorted(mots_majuscules.items(), key=lambda item: item[1], reverse=True)
        return mots_majuscules_tries

    def extraire_mot_cles(self, nom_fichier):
        mot_cles = []
        mots_et_occurrences = self.compter_mots_repetes(nom_fichier)
        mots_majuscules = self.extraire_mots_majuscules(nom_fichier)

        # Extraction des 5 premiers mots et leurs occurrences
        print("Les 5 premiers mots et leurs occurrences :")
        for word, count in mots_et_occurrences[:5]:
            mot_cles.append(word)
            print(f"{word}: {count} fois")

        print("Les 5 premiers mots en majuscules et leurs occurrences :")
        for word, count in mots_majuscules[:5]:
            mot_cles.append(word)
            print(f"{word}: {count} fois")

        return mot_cles
   